"""Rendering the file-shaped artifacts, and where they are kept.

The Studio's Summary and FAQ are text, so a note holds them. An infographic and a slide
deck are files, and a file needs somewhere to live and a URL to be fetched from — that is
all this module is: renderers, plus a small directory with opaque ids.

Both renderers take a filled Pydantic object and write bytes. No model is involved here,
which is the payoff of structured output: the layout is a function of the data.
"""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from core.store import DATA_DIR

ARTIFACT_DIR = DATA_DIR / "artifacts"
ID_PATTERN = re.compile(r"^[0-9a-f]{16}\.(html|pptx)$")

MEDIA_TYPES = {
    ".html": "text/html",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


@dataclass
class StoredFile:
    """A rendered artifact on disk."""

    file_id: str  # "<hex>.<ext>" — the whole name, so the extension survives the round trip
    path: Path
    download_name: str

    @property
    def url(self) -> str:
        return f"/api/studio/download/{self.file_id}"


def _safe_name(title: str, extension: str) -> str:
    """A filename the browser will accept, derived from the artifact's own title."""
    stem = re.sub(r"[^\w\s-]", "", title).strip() or "artifact"
    return re.sub(r"\s+", "-", stem)[:60].lower() + extension


def store_file(data: bytes, title: str, extension: str) -> StoredFile:
    """Write a rendered artifact and hand back the id it can be fetched by."""
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    file_id = f"{uuid.uuid4().hex[:16]}{extension}"
    path = ARTIFACT_DIR / file_id
    path.write_bytes(data)
    return StoredFile(file_id=file_id, path=path, download_name=_safe_name(title, extension))


def resolve(file_id: str) -> Path | None:
    """The file behind an id, or None.

    The id is matched against a strict pattern rather than sanitised: it is generated
    here, so anything that does not look exactly like one is not worth interpreting.
    """
    if not ID_PATTERN.match(file_id):
        return None
    path = ARTIFACT_DIR / file_id
    return path if path.is_file() else None


# -- infographic ---------------------------------------------------------------

_CSS = """
:root { color-scheme: light; }
* { box-sizing: border-box; }
body { margin: 0; padding: 48px 32px; background: #f5f7fb; color: #101828;
       font: 16px/1.6 -apple-system, "Segoe UI", Roboto, sans-serif; }
.sheet { max-width: 860px; margin: 0 auto; background: #fff; border-radius: 18px;
         padding: 40px; box-shadow: 0 18px 50px rgba(16,24,40,.10); }
h1 { margin: 0 0 6px; font-size: 30px; line-height: 1.25; }
.subtitle { margin: 0 0 32px; color: #667085; font-size: 17px; }
.stats { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
         gap: 14px; margin-bottom: 34px; }
.stat { background: #f0f5ff; border: 1px solid #d6e4ff; border-radius: 14px; padding: 18px; }
.stat .value { font-size: 28px; font-weight: 700; color: #1d4ed8; line-height: 1.1; }
.stat .label { font-size: 13px; margin-top: 6px; }
.stat .src { font-size: 11px; color: #667085; margin-top: 8px; }
section { margin-bottom: 26px; }
section h2 { font-size: 17px; margin: 0 0 10px; padding-bottom: 6px;
             border-bottom: 1px solid #eaecf0; }
section ul { margin: 0; padding-left: 20px; }
section li { margin-bottom: 6px; }
.takeaway { background: #101828; color: #fff; border-radius: 14px; padding: 22px; }
.takeaway .label { font-size: 11px; letter-spacing: .08em; text-transform: uppercase;
                   color: #98a2b3; margin-bottom: 6px; }
footer { margin-top: 26px; font-size: 12px; color: #98a2b3; text-align: center; }
@media print { body { background: #fff; padding: 0; } .sheet { box-shadow: none; } }
"""


def render_infographic(data) -> bytes:
    """A self-contained HTML page — no build step, no external assets, prints cleanly."""
    stats = "".join(
        f'<div class="stat"><div class="value">{escape(s.value)}</div>'
        f'<div class="label">{escape(s.label)}</div>'
        f'<div class="src">{escape(s.source)}</div></div>'
        for s in data.stats
    )
    sections = "".join(
        f"<section><h2>{escape(s.heading)}</h2><ul>"
        + "".join(f"<li>{escape(p)}</li>" for p in s.points)
        + "</ul></section>"
        for s in data.sections
    )
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(data.title)}</title><style>{_CSS}</style></head>
<body><div class="sheet">
<h1>{escape(data.title)}</h1>
<p class="subtitle">{escape(data.subtitle)}</p>
{f'<div class="stats">{stats}</div>' if stats else ""}
{sections}
<div class="takeaway"><div class="label">Takeaway</div>{escape(data.takeaway)}</div>
<footer>Generated from the sources in your notebook.</footer>
</div></body></html>""".encode()


# -- slide deck ----------------------------------------------------------------


def render_deck(data) -> bytes:
    """A .pptx built from the deck schema, one slide per entry."""
    presentation = Presentation()

    title_layout, bullet_layout = presentation.slide_layouts[0], presentation.slide_layouts[1]

    opening = presentation.slides.add_slide(title_layout)
    opening.shapes.title.text = data.title
    opening.placeholders[1].text = data.subtitle

    for slide_data in data.slides:
        slide = presentation.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_data.title

        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide_data.bullets):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0

        # the attribution goes in its own small box, not in the bullets
        note = slide.shapes.add_textbox(Inches(0.6), Inches(6.4), Inches(9), Inches(0.4))
        run = note.text_frame.paragraphs[0].add_run()
        run.text = f"Source: {slide_data.source}"
        run.font.size = Pt(11)

    from io import BytesIO

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
